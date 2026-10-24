"""Markdown, DOCX and PDF exports of already written reports."""

from __future__ import annotations

from pathlib import Path

FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def assert_inside(path: Path, root: Path) -> Path:
    resolved = path.resolve()
    root_resolved = root.resolve()
    if resolved != root_resolved and root_resolved not in resolved.parents:
        raise ValueError("export path is outside the storage root")
    return resolved


def write_markdown(markdown: str, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(markdown, encoding="utf-8")


def write_docx(markdown: str, path: Path) -> None:
    from docx import Document

    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        else:
            document.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_pdf(markdown: str, path: Path) -> None:
    from fpdf import FPDF

    if not Path(FONT).exists():
        raise FileNotFoundError("DejaVu font is not installed; PDF export is unavailable.")
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("DejaVu", "", FONT)
    pdf.add_font("DejaVu", "B", FONT_BOLD)
    pdf.add_page()
    pdf.set_font("DejaVu", size=11)
    for line in markdown.splitlines():
        if line.startswith("# "):
            pdf.set_font("DejaVu", "B", 16)
            pdf.multi_cell(0, 9, line[2:])
            pdf.set_font("DejaVu", size=11)
        elif line.startswith("## "):
            pdf.set_font("DejaVu", "B", 13)
            pdf.multi_cell(0, 8, line[3:])
            pdf.set_font("DejaVu", size=11)
        else:
            pdf.multi_cell(0, 6, line if line else " ")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(path)
